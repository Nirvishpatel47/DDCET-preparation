import os
import json
import csv
import io
import uuid
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
import sqlite3

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None

from database.db import get_db, BASE_DIR
from models.schemas import QuestionModel, ImportResult
from services.validation_service import validate_and_prepare_question, get_or_create_hierarchy
from services.duplicate_service import compute_question_hash, compute_question_text_hash
from services.markdown_parser import parse_markdown_questions

TEMPLATES_DIR = os.path.join(BASE_DIR, "import_templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)


SAMPLE_DEMO_QUESTIONS = [
    {
        "id": "DEMO-SAMPLE-MATH-001",
        "question": "What is the derivative of sin(2x) with respect to x?",
        "options": {
            "A": "2 * cos(2x)",
            "B": "-2 * cos(2x)",
            "C": "cos(2x)",
            "D": "2 * sin(2x)"
        },
        "correct_answer": "A",
        "explanation": "Using the chain rule: d/dx[sin(u)] = cos(u) * du/dx. Here u = 2x, du/dx = 2, so the derivative is 2*cos(2x).",
        "subject": "Mathematics",
        "topic": "Calculus",
        "subtopic": "Derivatives",
        "difficulty": "easy",
        "source": {
            "name": "DDCET Sample Demonstration Bank",
            "page": 1
        },
        "tags": ["calculus", "chain_rule", "sample_data"],
        "verified": True,
        "created_at": datetime.utcnow().isoformat()
    },
    {
        "id": "DEMO-SAMPLE-MATH-002",
        "question": "If matrix A is of order 2x3 and matrix B is of order 3x2, what is the order of product matrix AB?",
        "options": {
            "A": "3x3",
            "B": "2x2",
            "C": "2x3",
            "D": "Product is not defined"
        },
        "correct_answer": "B",
        "explanation": "For matrix multiplication A(m x n) and B(n x p), the resulting matrix has dimensions m x p. Here m=2 and p=2, so AB is a 2x2 matrix.",
        "subject": "Mathematics",
        "topic": "Matrices and Determinants",
        "subtopic": "Matrix Multiplication",
        "difficulty": "medium",
        "source": {
            "name": "DDCET Sample Demonstration Bank",
            "page": 2
        },
        "tags": ["matrices", "linear_algebra", "sample_data"],
        "verified": True,
        "created_at": datetime.utcnow().isoformat()
    },
    {
        "id": "DEMO-SAMPLE-ENG-001",
        "question": "According to Ohm's Law, the current flowing through a conductor is directly proportional to the voltage across it, provided which condition remains constant?",
        "options": {
            "A": "Electric charge",
            "B": "Temperature and physical dimensions",
            "C": "Magnetic flux",
            "D": "Capacitance"
        },
        "correct_answer": "B",
        "explanation": "Ohm's Law states I = V/R, which holds true provided temperature, pressure, and other physical conditions remain constant.",
        "subject": "Engineering Science",
        "topic": "Basic Electrical",
        "subtopic": "Ohm's Law",
        "difficulty": "easy",
        "source": {
            "name": "DDCET Sample Demonstration Bank",
            "page": 5
        },
        "tags": ["electrical", "ohms_law", "sample_data"],
        "verified": True,
        "created_at": datetime.utcnow().isoformat()
    }
]


def ensure_templates_exist():
    """Generates Excel, CSV, and JSON templates with clear instructions and sample demo rows."""
    # 1. JSON Template
    json_path = os.path.join(TEMPLATES_DIR, "ddcet_question_template.json")
    if not os.path.exists(json_path):
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(SAMPLE_DEMO_QUESTIONS, f, indent=2)

    # 2. CSV Template
    csv_path = os.path.join(TEMPLATES_DIR, "ddcet_question_template.csv")
    if not os.path.exists(csv_path):
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([
                "id", "question", "option_a", "option_b", "option_c", "option_d",
                "correct_answer", "explanation", "subject", "topic", "subtopic",
                "difficulty", "source_name", "source_page", "tags", "verified"
            ])
            for q in SAMPLE_DEMO_QUESTIONS:
                writer.writerow([
                    q["id"],
                    q["question"],
                    q["options"]["A"],
                    q["options"]["B"],
                    q["options"]["C"],
                    q["options"]["D"],
                    q["correct_answer"],
                    q["explanation"],
                    q["subject"],
                    q["topic"],
                    q["subtopic"],
                    q["difficulty"],
                    q["source"]["name"],
                    q["source"]["page"],
                    ",".join(q["tags"]),
                    "true" if q["verified"] else "false"
                ])

    # 3. Excel Template (.xlsx)
    xlsx_path = os.path.join(TEMPLATES_DIR, "ddcet_question_template.xlsx")
    if openpyxl and not os.path.exists(xlsx_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DDCET Questions Template"

        headers = [
            "id", "question", "option_a", "option_b", "option_c", "option_d",
            "correct_answer", "explanation", "subject", "topic", "subtopic",
            "difficulty", "source_name", "source_page", "tags", "verified"
        ]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for q in SAMPLE_DEMO_QUESTIONS:
            ws.append([
                q["id"],
                q["question"],
                q["options"]["A"],
                q["options"]["B"],
                q["options"]["C"],
                q["options"]["D"],
                q["correct_answer"],
                q["explanation"],
                q["subject"],
                q["topic"],
                q["subtopic"],
                q["difficulty"],
                q["source"]["name"],
                q["source"]["page"],
                ",".join(q["tags"]),
                "TRUE" if q["verified"] else "FALSE"
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(xlsx_path)

    # 4. Markdown Template (.md) with Mathematical Formulas
    md_path = os.path.join(TEMPLATES_DIR, "ddcet_question_template.md")
    if not os.path.exists(md_path):
        sample_md_content = """# Gujarat DDCET Mathematics Question Bank Template (.md)

---
## Question: DDCET-MATH-001
**Subject**: Mathematics
**Topic**: Calculus
**Subtopic**: Derivatives
**Difficulty**: easy
**Verified**: true

What is the derivative of $\\sin(2x)$ with respect to $x$?

- (A) $2\\cos(2x)$
- (B) $-2\\cos(2x)$
- (C) $\\cos(2x)$
- (D) $2\\sin(2x)$

**Correct Answer**: A
**Explanation**:
Using the chain rule of differentiation:
$$\\frac{d}{dx}[\\sin(u)] = \\cos(u) \\cdot \\frac{du}{dx}$$
Here $u = 2x$ and $\\frac{du}{dx} = 2$.
Therefore, $\\frac{d}{dx}[\\sin(2x)] = 2\\cos(2x)$.

---
## Question: DDCET-MATH-002
**Subject**: Mathematics
**Topic**: Calculus
**Subtopic**: Definite Integrals
**Difficulty**: medium
**Verified**: true

Evaluate the definite integral:
$$\\int_0^{\\frac{\\pi}{2}} \\sin^2(x) \\, dx$$

- (A) $\\frac{\\pi}{4}$
- (B) $\\frac{\\pi}{2}$
- (C) $1$
- (D) $0$

**Correct Answer**: A
**Explanation**:
Using the half-angle identity $\\sin^2(x) = \\frac{1 - \\cos(2x)}{2}$:
$$\\int_0^{\\frac{\\pi}{2}} \\frac{1 - \\cos(2x)}{2} \\, dx = \\left[ \\frac{x}{2} - \\frac{\\sin(2x)}{4} \\right]_0^{\\frac{\\pi}{2}} = \\frac{\\pi}{4}$$

---
## Question: DDCET-MATH-003
**Subject**: Mathematics
**Topic**: Matrices and Determinants
**Subtopic**: Eigenvalues
**Difficulty**: medium
**Verified**: true

If matrix $A$ is given by:
$$A = \\begin{pmatrix} 2 & 0 \\\\ 0 & 5 \\end{pmatrix}$$
What are the eigenvalues of $A$?

- (A) $\\lambda = 2, 5$
- (B) $\\lambda = 0, 7$
- (C) $\\lambda = 10, 1$
- (D) $\\lambda = -2, -5$

**Correct Answer**: A
**Explanation**:
For any diagonal matrix, the eigenvalues are simply the diagonal entries:
$$\\det(A - \\lambda I) = (2 - \\lambda)(5 - \\lambda) = 0 \\implies \\lambda = 2, 5$$
"""
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(sample_md_content)


# Ensure templates are ready at import time
ensure_templates_exist()


def parse_raw_file(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parses raw uploaded bytes from JSON, CSV, XLSX, or Markdown (.md) into a list of dictionaries."""
    lower_name = filename.lower()
    records = []

    if lower_name.endswith(".md") or lower_name.endswith(".markdown"):
        content = file_bytes.decode("utf-8")
        records = parse_markdown_questions(content)

    elif lower_name.endswith(".json"):
        content = file_bytes.decode("utf-8")
        parsed = json.loads(content)
        if isinstance(parsed, list):
            records = parsed
        elif isinstance(parsed, dict) and "questions" in parsed:
            records = parsed["questions"]
        else:
            records = [parsed]

    elif lower_name.endswith(".csv"):
        content = file_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(content))
        for row in reader:
            records.append(dict(row))

    elif lower_name.endswith(".xlsx"):
        if not openpyxl:
            raise RuntimeError("openpyxl is required to parse Excel files.")
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
        for r in rows[1:]:
            if all(v is None for v in r):
                continue
            item = {}
            for h, val in zip(header, r):
                item[h] = val
            records.append(item)
    else:
        raise ValueError("Unsupported file format. Supported formats: .md, .json, .csv, .xlsx")

    return records


def process_import(
    file_bytes: bytes,
    filename: str,
    preview_only: bool = False
) -> Dict[str, Any]:
    """
    Two-stage import engine:
    1. Parse and validate structure
    2. Check duplicates (exact and probable)
    3. If preview_only: return inspection report without DB commit
    4. If commit: insert valid questions, quarantine rejected/invalid rows
    """
    raw_records = parse_raw_file(file_bytes, filename)
    batch_id = str(uuid.uuid4())

    valid_models: List[Tuple[QuestionModel, str, str]] = [] # (model, text_hash, q_text_hash)
    rejected_records: List[Dict[str, Any]] = []
    exact_duplicate_records: List[Dict[str, Any]] = []
    probable_duplicates: List[Dict[str, Any]] = []

    seen_in_batch_hashes = set()
    seen_in_batch_q_hashes = set()
    seen_in_batch_ids = set()

    with get_db() as conn:
        for idx, raw in enumerate(raw_records):
            row_num = idx + 1
            # Check batch-internal ID duplicate
            qid = str(raw.get("id", "")).strip()
            if qid and qid in seen_in_batch_ids:
                rejected_records.append({
                    "row": row_num,
                    "id": qid,
                    "error": f"Duplicate question ID '{qid}' found within the import file itself.",
                    "data": raw
                })
                continue
            if qid:
                seen_in_batch_ids.add(qid)

            model, error_msg, exact_dup, prob_dups = validate_and_prepare_question(raw, conn)
            
            if error_msg:
                if exact_dup:
                    exact_duplicate_records.append({
                        "row": row_num,
                        "id": raw.get("id"),
                        "question": raw.get("question"),
                        "match_existing_id": exact_dup.get("existing_id") or exact_dup.get("id"),
                        "reason": error_msg,
                        "data": raw
                    })
                else:
                    rejected_records.append({
                        "row": row_num,
                        "id": raw.get("id"),
                        "error": error_msg,
                        "data": raw
                    })
                continue

            text_hash = compute_question_hash(model.question, model.options)
            q_text_hash = compute_question_text_hash(model.question)

            if text_hash in seen_in_batch_hashes:
                exact_duplicate_records.append({
                    "row": row_num,
                    "id": model.id,
                    "question": model.question,
                    "reason": "Exact duplicate of another question within this same import file.",
                    "data": raw
                })
                continue

            if q_text_hash in seen_in_batch_q_hashes:
                exact_duplicate_records.append({
                    "row": row_num,
                    "id": model.id,
                    "question": model.question,
                    "reason": "Question text already appears earlier within this same import file.",
                    "data": raw
                })
                continue

            seen_in_batch_hashes.add(text_hash)
            seen_in_batch_q_hashes.add(q_text_hash)

            if prob_dups:
                probable_duplicates.append({
                    "row": row_num,
                    "id": model.id,
                    "question": model.question,
                    "candidates": prob_dups
                })

            valid_models.append((model, text_hash, q_text_hash))

        # If not preview, commit to database
        if not preview_only:
            now_iso = datetime.utcnow().isoformat()
            for model, thash, qthash in valid_models:
                sub_id, top_id, subtop_id = get_or_create_hierarchy(
                    conn, model.subject, model.topic, model.subtopic
                )
                source_json = json.dumps(model.source.model_dump() if model.source else {})
                tags_json = json.dumps(model.tags)
                verified_int = 1 if model.verified else 0

                conn.execute(
                    """
                    INSERT INTO questions (
                        id, question_text, option_a, option_b, option_c, option_d,
                        correct_answer, explanation, subject_id, topic_id, subtopic_id,
                        difficulty, source_json, tags_json, verified, text_hash, question_text_hash, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        model.id,
                        model.question,
                        model.options["A"],
                        model.options["B"],
                        model.options["C"],
                        model.options["D"],
                        model.correct_answer,
                        model.explanation,
                        sub_id,
                        top_id,
                        subtop_id,
                        model.difficulty,
                        source_json,
                        tags_json,
                        verified_int,
                        thash,
                        qthash,
                        model.created_at
                    )
                )

            # Store rejected records in quarantine table for error review
            for rej in rejected_records:
                conn.execute(
                    """
                    INSERT INTO quarantined_imports (batch_id, raw_data_json, reason, quarantined_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (batch_id, json.dumps(rej["data"]), rej["error"], now_iso)
                )

    return {
        "batch_id": batch_id,
        "total_read": len(raw_records),
        "imported": len(valid_models) if not preview_only else 0,
        "valid_count": len(valid_models),
        "rejected_count": len(rejected_records),
        "exact_duplicates_count": len(exact_duplicate_records),
        "probable_duplicates_count": len(probable_duplicates),
        "preview_only": preview_only,
        "sample_valid": [m[0].model_dump() for m in valid_models[:5]],
        "rejected": rejected_records,
        "exact_duplicates": exact_duplicate_records,
        "probable_duplicates": probable_duplicates
    }


def export_questions(
    subject_id: Optional[int] = None,
    topic_id: Optional[int] = None,
    verified_only: bool = False,
    format_type: str = "json"
) -> Tuple[bytes, str, str]:
    """Exports questions in JSON, CSV, or XLSX format."""
    with get_db() as conn:
        query = """
            SELECT q.id, q.question_text, q.option_a, q.option_b, q.option_c, q.option_d,
                   q.correct_answer, q.explanation, q.difficulty, q.source_json, q.tags_json,
                   q.verified, q.created_at, s.name as subject, t.name as topic, st.name as subtopic
            FROM questions q
            JOIN subjects s ON q.subject_id = s.id
            JOIN topics t ON q.topic_id = t.id
            JOIN subtopics st ON q.subtopic_id = st.id
            WHERE 1=1
        """
        params = []
        if subject_id:
            query += " AND q.subject_id = ?"
            params.append(subject_id)
        if topic_id:
            query += " AND q.topic_id = ?"
            params.append(topic_id)
        if verified_only:
            query += " AND q.verified = 1"
        query += " ORDER BY q.subject_id, q.topic_id, q.id ASC"

        rows = conn.execute(query, params).fetchall()

    export_list = []
    for r in rows:
        source_data = json.loads(r["source_json"]) if r["source_json"] else {}
        tags_data = json.loads(r["tags_json"]) if r["tags_json"] else []
        export_list.append({
            "id": r["id"],
            "question": r["question_text"],
            "options": {
                "A": r["option_a"],
                "B": r["option_b"],
                "C": r["option_c"],
                "D": r["option_d"]
            },
            "correct_answer": r["correct_answer"],
            "explanation": r["explanation"],
            "subject": r["subject"],
            "topic": r["topic"],
            "subtopic": r["subtopic"],
            "difficulty": r["difficulty"],
            "source": source_data,
            "tags": tags_data,
            "verified": bool(r["verified"]),
            "created_at": r["created_at"]
        })

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if format_type == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "id", "question", "option_a", "option_b", "option_c", "option_d",
            "correct_answer", "explanation", "subject", "topic", "subtopic",
            "difficulty", "source_name", "source_page", "tags", "verified", "created_at"
        ])
        for q in export_list:
            writer.writerow([
                q["id"],
                q["question"],
                q["options"]["A"],
                q["options"]["B"],
                q["options"]["C"],
                q["options"]["D"],
                q["correct_answer"],
                q["explanation"] or "",
                q["subject"],
                q["topic"],
                q["subtopic"],
                q["difficulty"],
                q["source"].get("name", ""),
                q["source"].get("page", ""),
                ",".join(q["tags"]),
                "true" if q["verified"] else "false",
                q["created_at"]
            ])
        return output.getvalue().encode("utf-8"), f"ddcet_questions_{timestamp}.csv", "text/csv"

    elif format_type == "xlsx":
        if not openpyxl:
            raise RuntimeError("openpyxl required for Excel export")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DDCET Questions"
        headers = [
            "id", "question", "option_a", "option_b", "option_c", "option_d",
            "correct_answer", "explanation", "subject", "topic", "subtopic",
            "difficulty", "source_name", "source_page", "tags", "verified", "created_at"
        ]
        ws.append(headers)
        for q in export_list:
            ws.append([
                q["id"],
                q["question"],
                q["options"]["A"],
                q["options"]["B"],
                q["options"]["C"],
                q["options"]["D"],
                q["correct_answer"],
                q["explanation"] or "",
                q["subject"],
                q["topic"],
                q["subtopic"],
                q["difficulty"],
                q["source"].get("name", ""),
                q["source"].get("page", ""),
                ",".join(q["tags"]),
                "TRUE" if q["verified"] else "FALSE",
                q["created_at"]
            ])
        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue(), f"ddcet_questions_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    elif format_type == "md" or format_type == "markdown":
        md_lines = ["# Gujarat DDCET Question Bank Export\n"]
        for q in export_list:
            md_lines.append(f"---\n## Question: {q['id']}")
            md_lines.append(f"**Subject**: {q['subject']}")
            md_lines.append(f"**Topic**: {q['topic']}")
            md_lines.append(f"**Subtopic**: {q['subtopic']}")
            md_lines.append(f"**Difficulty**: {q['difficulty']}")
            md_lines.append(f"**Verified**: {'true' if q['verified'] else 'false'}\n")
            md_lines.append(f"{q['question']}\n")
            md_lines.append(f"- (A) {q['options']['A']}")
            md_lines.append(f"- (B) {q['options']['B']}")
            md_lines.append(f"- (C) {q['options']['C']}")
            md_lines.append(f"- (D) {q['options']['D']}\n")
            md_lines.append(f"**Correct Answer**: {q['correct_answer']}")
            if q["explanation"]:
                md_lines.append(f"**Explanation**:\n{q['explanation']}\n")
            else:
                md_lines.append("")
        md_text = "\n".join(md_lines)
        return md_text.encode("utf-8"), f"ddcet_questions_{timestamp}.md", "text/markdown"

    else:
        # Default JSON
        json_data = json.dumps(export_list, indent=2).encode("utf-8")
        return json_data, f"ddcet_questions_{timestamp}.json", "application/json"


def export_quarantined_records(batch_id: Optional[str] = None) -> List[Dict[str, Any]]:
    """Retrieves quarantined rejected records for correction."""
    with get_db() as conn:
        query = "SELECT id, batch_id, raw_data_json, reason, quarantined_at FROM quarantined_imports"
        params = []
        if batch_id:
            query += " WHERE batch_id = ?"
            params.append(batch_id)
        query += " ORDER BY id DESC"
        rows = conn.execute(query, params).fetchall()
        
        results = []
        for r in rows:
            results.append({
                "quarantine_id": r["id"],
                "batch_id": r["batch_id"],
                "reason": r["reason"],
                "quarantined_at": r["quarantined_at"],
                "data": json.loads(r["raw_data_json"])
            })
        return results
